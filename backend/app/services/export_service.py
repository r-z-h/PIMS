"""Excel 导出：单个文件 11 个工作表（docs/01-开发手册.md 5.10、docs/03-接口设计.md 第 11 节）。

Sheet 顺序：
项目盈亏 / 资金流水 / 人工 / 材料 / 机械 / 杂费 / 垫付 / 应收 / 应付 / 库存 / 人工结算

格式要求：中文表头加粗、冻结首行、列宽自适应、金额千分位两位小数。
"""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models.ar_ap import TYPE_PAYABLE, TYPE_RECEIVABLE
from ..models.cost import (
    TYPE_ADVANCE,
    TYPE_LABOR,
    TYPE_MACHINE,
    TYPE_MISC,
    TYPE_MATERIAL,
)
from ..models.funds_flow import BIZ_TYPE_LABELS, DIRECTION_LABELS, PAY_METHOD_LABELS
from ..models.machine_item import MEASURE_LABELS
from ..models.project import STATUS_ONGOING, Project
from ..utils.dates import check_range
from ..utils.money import ZERO, money, qty, qty2
from . import (
    arap_service,
    cost_service,
    dashboard_service,
    funds_service,
    inventory_service,
    labor_service,
)

SHEET_NAMES = [
    "项目盈亏",
    "资金流水",
    "人工",
    "材料",
    "机械",
    "杂费",
    "垫付",
    "应收",
    "应付",
    "库存",
    "人工结算",
]

MONEY_FMT = "#,##0.00"
QTY3_FMT = "#,##0.000"
QTY2_FMT = "#,##0.00"
DATE_FMT = "@"

HEADER_FILL = PatternFill("solid", fgColor="F0F0F0")
MAX_ROWS = 100_000


def _display_width(value: Any, number_format: str | None = None) -> int:
    if value is None:
        return 0
    if isinstance(value, Decimal):
        text = f"{value:,.2f}" if number_format == MONEY_FMT else format(value, "f")
    elif isinstance(value, (datetime, date)):
        text = value.strftime("%Y-%m-%d")
    else:
        text = str(value)
    # 中文按 2 个字符宽度估算
    return sum(2 if ord(ch) > 127 else 1 for ch in text)


def _append_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
    formats: dict[int, str] | None = None,
) -> Worksheet:
    formats = formats or {}
    sheet = workbook.create_sheet(title=title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = HEADER_FILL
    for row in rows:
        sheet.append(row)

    # 金额 / 数量列格式
    for column_index, number_format in formats.items():
        for row_index in range(2, sheet.max_row + 1):
            sheet.cell(row=row_index, column=column_index).number_format = number_format

    # 冻结首行
    sheet.freeze_panes = "A2"

    # 列宽自适应
    for column_index, header in enumerate(headers, start=1):
        width = _display_width(header)
        for row_index in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row_index, column=column_index)
            width = max(width, _display_width(cell.value, formats.get(column_index)))
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(width + 3, 9), 42)
    return sheet


def _status_label(status: str) -> str:
    return "在建" if status == STATUS_ONGOING else "完工"


def _project_names(db: Session) -> dict[int, str]:
    return {pid: name for pid, name in db.execute(select(Project.id, Project.name)).all()}


def _profit_rows(db: Session, project_id: int | None, start: date | None, end: date | None) -> list[list[Any]]:
    table = dashboard_service.project_table(db, start, end)
    items = table["items"]
    if project_id is not None:
        items = [item for item in items if item["project_id"] == int(project_id)]
    rows = [
        [
            item["name"],
            _status_label(item["status"]),
            money(item["client_budget"]),
            money(item["cost_budget"]),
            money(item["expect_profit"]),
            money(item["actual_cost"]),
            money(item["settle_profit"]),
            money(item["budget_left"]),
            money(item["received"]),
            money(item["unreceived"]),
            money(item["paid"]),
            money(item["unpaid"]),
        ]
        for item in items
    ]
    total = [sum((row[index] for row in rows), ZERO) for index in range(2, 12)]
    if rows:
        rows.append(["合计", "", *[money(value) for value in total]])
    return rows


def _funds_rows(db: Session, project_id: int | None, start: date | None, end: date | None) -> list[list[Any]]:
    items, _ = funds_service.list_flows(
        db, project_id=project_id, start=start, end=end, page=1, page_size=MAX_ROWS
    )
    return [
        [
            row["biz_date"],
            row.get("project_name"),
            DIRECTION_LABELS.get(row["direction"], row["direction"]),
            BIZ_TYPE_LABELS.get(row["biz_type"], row["biz_type"]),
            money(row["amount"]),
            row["partner"],
            PAY_METHOD_LABELS.get(row["pay_method"], "") if row["pay_method"] else "",
            row["remark"],
        ]
        for row in items
    ]


def _cost_rows(
    db: Session, cost_type: str, project_id: int | None, start: date | None, end: date | None
) -> list[dict[str, Any]]:
    items, _ = cost_service.list_costs(
        db,
        types=[cost_type],
        project_id=project_id,
        start=start,
        end=end,
        page=1,
        page_size=MAX_ROWS,
    )
    return items


def _labor_rows(db: Session, project_id: int | None, start: date | None, end: date | None) -> list[list[Any]]:
    return [
        [
            row["biz_date"],
            row.get("project_name"),
            row.get("worker_name"),
            qty2(row.get("work_days") or 0),
            money(row.get("day_rate") or 0),
            money(row["amount"]),
            row["remark"],
        ]
        for row in _cost_rows(db, TYPE_LABOR, project_id, start, end)
    ]


def _material_rows(db: Session, project_id: int | None, start: date | None, end: date | None) -> list[list[Any]]:
    return [
        [
            row["biz_date"],
            row.get("project_name"),
            row.get("material_name"),
            qty(row.get("quantity") or 0),
            row.get("unit"),
            money(row.get("unit_price") or 0),
            money(row["amount"]),
            row.get("supplier"),
            row["remark"],
        ]
        for row in _cost_rows(db, TYPE_MATERIAL, project_id, start, end)
    ]


def _machine_rows(db: Session, project_id: int | None, start: date | None, end: date | None) -> list[list[Any]]:
    return [
        [
            row["biz_date"],
            row.get("project_name"),
            row.get("machine_name"),
            MEASURE_LABELS.get(row.get("measure"), row.get("measure")),
            qty2(row.get("quantity") or 0),
            money(row.get("unit_price") or 0),
            money(row["amount"]),
            row.get("lessor"),
            row["remark"],
        ]
        for row in _cost_rows(db, TYPE_MACHINE, project_id, start, end)
    ]


def _misc_rows(db: Session, project_id: int | None, start: date | None, end: date | None) -> list[list[Any]]:
    return [
        [
            row["biz_date"],
            row.get("project_name"),
            row.get("category"),
            money(row["amount"]),
            row["remark"],
        ]
        for row in _cost_rows(db, TYPE_MISC, project_id, start, end)
    ]


def _advance_rows(db: Session, project_id: int | None, start: date | None, end: date | None) -> list[list[Any]]:
    return [
        [
            row["biz_date"],
            row.get("project_name"),
            row.get("payer"),
            money(row["amount"]),
            row.get("purpose"),
            row["remark"],
        ]
        for row in _cost_rows(db, TYPE_ADVANCE, project_id, start, end)
    ]


def _arap_rows(db: Session, type_: str, project_id: int | None, start: date | None, end: date | None) -> list[list[Any]]:
    items, _ = arap_service.list_arap(
        db, type_=type_, project_id=project_id, start=start, end=end, page=1, page_size=MAX_ROWS
    )
    return [
        [
            row["biz_date"],
            row.get("project_name"),
            row["biz_type_label"],
            row["partner"],
            money(row["amount"]),
            money(row["settled_amount"]),
            money(row["unsettled"]),
            row["status_label"],
        ]
        for row in items
    ]


def _inventory_rows(db: Session, project_id: int | None) -> list[list[Any]]:
    names = _project_names(db)
    return [
        [
            names.get(row["project_id"]),
            row["material_name"],
            row["unit"],
            qty(row["in_qty"]),
            qty(row["out_qty"]),
            qty(row["balance"]),
        ]
        for row in inventory_service.list_balances(db, project_id)
    ]


def _settle_rows(db: Session, project_id: int | None) -> list[list[Any]]:
    items, _ = labor_service.list_settles(db, project_id=project_id, page=1, page_size=MAX_ROWS)
    return [
        [
            row["project_name"],
            row["worker_name"],
            qty2(row["total_days"]),
            money(row["amount_due"]),
            money(row["paid_amount"]),
            money(row["unpaid_amount"]),
            row["status_label"],
            row["settle_date"],
            row["pay_date"],
            PAY_METHOD_LABELS.get(row["pay_method"], row["pay_method"]) if row["pay_method"] else "",
        ]
        for row in items
    ]


def build_workbook(
    db: Session,
    project_id: int | None = None,
    start: date | None = None,
    end: date | None = None,
) -> Workbook:
    """按筛选条件生成 11 个 Sheet 的工作簿。"""
    check_range(start, end)
    workbook = Workbook()
    workbook.remove(workbook.active)

    _append_sheet(
        workbook,
        "项目盈亏",
        [
            "项目",
            "状态",
            "甲方预算",
            "我方成本预算",
            "预计利润",
            "实际成本",
            "结算利润",
            "成本预算余额",
            "已收",
            "未收",
            "已付",
            "未付",
        ],
        _profit_rows(db, project_id, start, end),
        formats={index: MONEY_FMT for index in range(3, 13)},
    )
    _append_sheet(
        workbook,
        "资金流水",
        ["日期", "项目", "收付", "类型", "金额", "对方", "支付方式", "备注"],
        _funds_rows(db, project_id, start, end),
        formats={5: MONEY_FMT},
    )
    _append_sheet(
        workbook,
        "人工",
        ["日期", "项目", "姓名", "工日数", "日工价", "金额", "备注"],
        _labor_rows(db, project_id, start, end),
        formats={4: QTY2_FMT, 5: MONEY_FMT, 6: MONEY_FMT},
    )
    _append_sheet(
        workbook,
        "材料",
        ["日期", "项目", "材料", "数量", "单位", "单价", "金额", "供应商", "备注"],
        _material_rows(db, project_id, start, end),
        formats={4: QTY3_FMT, 6: MONEY_FMT, 7: MONEY_FMT},
    )
    _append_sheet(
        workbook,
        "机械",
        ["日期", "项目", "设备", "计量方式", "数量", "单价", "金额", "出租方", "备注"],
        _machine_rows(db, project_id, start, end),
        formats={5: QTY2_FMT, 6: MONEY_FMT, 7: MONEY_FMT},
    )
    _append_sheet(
        workbook,
        "杂费",
        ["日期", "项目", "小类", "金额", "备注"],
        _misc_rows(db, project_id, start, end),
        formats={4: MONEY_FMT},
    )
    _append_sheet(
        workbook,
        "垫付",
        ["日期", "项目", "垫付人", "金额", "事由", "备注"],
        _advance_rows(db, project_id, start, end),
        formats={4: MONEY_FMT},
    )
    _append_sheet(
        workbook,
        "应收",
        ["日期", "项目", "类型", "对方", "金额", "已冲销", "未结", "状态"],
        _arap_rows(db, TYPE_RECEIVABLE, project_id, start, end),
        formats={5: MONEY_FMT, 6: MONEY_FMT, 7: MONEY_FMT},
    )
    _append_sheet(
        workbook,
        "应付",
        ["日期", "项目", "类型", "对方", "金额", "已冲销", "未结", "状态"],
        _arap_rows(db, TYPE_PAYABLE, project_id, start, end),
        formats={5: MONEY_FMT, 6: MONEY_FMT, 7: MONEY_FMT},
    )
    _append_sheet(
        workbook,
        "库存",
        ["项目", "材料", "单位", "入库合计", "领用合计", "余量"],
        _inventory_rows(db, project_id),
        formats={4: QTY3_FMT, 5: QTY3_FMT, 6: QTY3_FMT},
    )
    _append_sheet(
        workbook,
        "人工结算",
        ["项目", "姓名", "总工天", "应发金额", "已发金额", "未发金额", "状态", "结算日期", "发放日期", "发放方式"],
        _settle_rows(db, project_id),
        formats={3: QTY2_FMT, 4: MONEY_FMT, 5: MONEY_FMT, 6: MONEY_FMT},
    )
    return workbook


def export_excel(
    db: Session,
    project_id: int | None = None,
    start: date | None = None,
    end: date | None = None,
) -> tuple[bytes, str]:
    """返回 (xlsx 字节流, 建议文件名)。"""
    workbook = build_workbook(db, project_id, start, end)
    buffer = io.BytesIO()
    workbook.save(buffer)
    filename = f"pims-export-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
    return buffer.getvalue(), filename


__all__ = ["SHEET_NAMES", "build_workbook", "export_excel"]
