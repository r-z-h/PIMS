"""docs/01-开发手册.md 第 7 节 —— 验收标准 A1~A10 逐条可验。"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from helpers import SECURITY_ANSWERS, err_data, money_eq, ok_data

STATE: dict[str, int] = {}


def test_A1_用默认密码登录_要求设置密保(client):
    data = ok_data(client.post("/api/auth/login", json={"password": "123456"}))
    assert data["logged_in"] is True
    assert data["security_ready"] is False  # 前端跳「设置密保」
    assert data["is_initial_password"] is True


def test_A2_密保全对才能改密(client):
    ok_data(
        client.post(
            "/api/auth/security/setup",
            json={
                "questions": [
                    {"question": "问题一？", "answer": "答案一"},
                    {"question": "问题二？", "answer": "答案二"},
                    {"question": "问题三？", "answer": "答案三"},
                ]
            },
        )
    )
    # 答错任一题 → 拒绝
    err_data(
        client.post(
            "/api/auth/password/change",
            json={"new_password": "abc12345", "answers": ["答案一", "错的答案", "答案三"]},
        ),
        403,
    )
    # 三题全对 → 成功
    ok_data(
        client.post(
            "/api/auth/password/change",
            json={"new_password": "abc12345", "answers": ["答案一", "答案二", "答案三"]},
        )
    )
    ok_data(client.post("/api/auth/login", json={"password": "abc12345"}))


def test_A3_建项目400万_350万_预计利润50万(client):
    data = ok_data(
        client.post(
            "/api/projects",
            json={"name": "验收项目·某小区5#楼", "client_budget": 4000000, "cost_budget": 3500000},
        )
    )
    money_eq(data["expect_profit"], "500000.00")
    STATE["project_id"] = int(data["id"])

    table = ok_data(client.get("/api/dashboard/projects"))
    row = next(item for item in table["items"] if item["project_id"] == STATE["project_id"])
    money_eq(row["expect_profit"], "500000.00")


def test_A4_五类成本合计135_15万_结算利润264_85万(client):
    project_id = STATE["project_id"]
    ok_data(
        client.post(
            "/api/costs/labor",
            json={"project_id": project_id, "worker_name": "张三", "biz_date": "2026-05-01", "work_days": 5, "day_rate": 300},
        )
    )
    ok_data(
        client.post(
            "/api/costs/material",
            json={
                "project_id": project_id,
                "material_name": "水泥",
                "unit": "吨",
                "quantity": 2500,
                "unit_price": 400,
                "biz_date": "2026-05-02",
            },
        )
    )
    ok_data(
        client.post(
            "/api/costs/machine",
            json={
                "project_id": project_id,
                "machine_name": "挖机",
                "measure": "shift",
                "quantity": 100,
                "unit_price": 2000,
                "biz_date": "2026-05-03",
            },
        )
    )
    ok_data(
        client.post(
            "/api/costs/misc",
            json={"project_id": project_id, "category": "伙食费", "amount": 100000, "biz_date": "2026-05-04"},
        )
    )
    ok_data(
        client.post(
            "/api/costs/advance",
            json={"project_id": project_id, "payer": "老板", "purpose": "临时买料", "amount": 50000, "biz_date": "2026-05-05"},
        )
    )

    metrics = ok_data(api_overview(client, project_id))["metrics"]
    money_eq(metrics["actual_cost"], "1351500.00")  # 135.15 万
    money_eq(metrics["settle_profit"], "2648500.00")  # 400 − 135.15 = 264.85 万
    money_eq(metrics["expect_profit"], "500000.00")
    money_eq(metrics["budget_left"], "2148500.00")
    assert float(metrics["budget_usage_rate"]) > 38 and float(metrics["budget_usage_rate"]) < 39


def api_overview(client, project_id: int):
    return client.get(f"/api/projects/{project_id}/overview")


def test_A5_采购100吨领用80吨_余量20_成本只算采购(client):
    project_id = STATE["project_id"]
    before = ok_data(api_overview(client, project_id))["metrics"]["actual_cost"]

    ok_data(
        client.post(
            "/api/costs/material",
            json={
                "project_id": project_id,
                "material_name": "钢筋",
                "unit": "吨",
                "quantity": 100,
                "unit_price": 4000,
                "biz_date": "2026-05-06",
            },
        )
    )
    after_purchase = ok_data(api_overview(client, project_id))["metrics"]["actual_cost"]
    assert float(after_purchase) - float(before) == 400000.0

    out = ok_data(
        client.post(
            "/api/inventory/out",
            json={"project_id": project_id, "material_name": "钢筋", "unit": "吨", "quantity": 80, "biz_date": "2026-05-07"},
        )
    )
    assert str(out["balance"]) in ("20.000", "20.0", "20")
    after_out = ok_data(api_overview(client, project_id))["metrics"]["actual_cost"]
    assert after_out == after_purchase  # 领用不产生成本

    balances = ok_data(client.get("/api/inventory", params={"project_id": project_id}))["items"]
    steel = next(item for item in balances if item["material_name"] == "钢筋")
    assert str(steel["balance"]) in ("20.000", "20.0", "20")


def test_A6_应收50万收30万冲销_未收20万_结算利润不变(client):
    project_id = STATE["project_id"]
    settle_profit_before = ok_data(api_overview(client, project_id))["metrics"]["settle_profit"]

    arap_id = int(
        ok_data(
            client.post(
                "/api/arap",
                json={
                    "project_id": project_id,
                    "type": "receivable",
                    "biz_type": "progress",
                    "amount": 500000,
                    "biz_date": "2026-06-01",
                },
            )
        )["id"]
    )
    flow_id = int(
        ok_data(
            client.post(
                "/api/funds",
                json={
                    "project_id": project_id,
                    "direction": "in",
                    "biz_type": "progress",
                    "amount": 300000,
                    "biz_date": "2026-06-20",
                },
            )
        )["id"]
    )
    ok_data(
        client.post(
            "/api/arap/offset",
            json={"funds_flow_id": flow_id, "offsets": [{"ar_ap_id": arap_id, "amount": 300000}], "offset_date": "2026-06-20"},
        )
    )

    detail = ok_data(client.get(f"/api/arap/{arap_id}"))
    money_eq(detail["unsettled"], "200000.00")
    assert detail["status"] == "partial"

    metrics = ok_data(api_overview(client, project_id))["metrics"]
    assert metrics["settle_profit"] == settle_profit_before
    money_eq(metrics["actual_cost"], "1751500.00")


def test_A7_录工天后结算_结算金额等于已入成本且成本不变(client):
    project_id = STATE["project_id"]
    before = ok_data(api_overview(client, project_id))["metrics"]["actual_cost"]

    ok_data(
        client.post(
            "/api/costs/labor",
            json={"project_id": project_id, "worker_name": "李四", "biz_date": "2026-05-08", "work_days": 10, "day_rate": 350},
        )
    )
    after_labor = ok_data(api_overview(client, project_id))["metrics"]["actual_cost"]
    assert float(after_labor) - float(before) == 3500.0

    preview = ok_data(client.get("/api/labor/preview", params={"project_id": project_id}))
    li = next(item for item in preview["items"] if item["worker_name"] == "李四")
    money_eq(li["amount_due"], "3500.00")

    settled = ok_data(client.post("/api/labor/settle", json={"project_id": project_id, "settle_date": "2026-07-01"}))
    assert settled["created"] >= 1
    settle_ids = {item["worker_name"]: item["id"] for item in settled["items"]}
    ok_data(
        client.post(
            f"/api/labor/settles/{settle_ids['李四']}/pay",
            json={"amount": 3500, "pay_date": "2026-07-02", "pay_method": "bank_card"},
        )
    )
    after_settle = ok_data(api_overview(client, project_id))["metrics"]["actual_cost"]
    assert after_settle == after_labor  # 结算与发放不重复入成本
    STATE["labor_cost_id"] = int(
        ok_data(client.get("/api/costs", params={"project_id": project_id, "type": "labor"}))["items"][0]["id"]
    )


def test_A8_删除成本后实际成本减少且变更日志有记录(client):
    project_id = STATE["project_id"]
    before = float(ok_data(api_overview(client, project_id))["metrics"]["actual_cost"])
    cost_id = STATE["labor_cost_id"]

    ok_data(client.delete(f"/api/costs/{cost_id}"))
    after = float(ok_data(api_overview(client, project_id))["metrics"]["actual_cost"])
    assert after < before

    costs = ok_data(client.get("/api/costs", params={"project_id": project_id}))["items"]
    assert all(item["id"] != cost_id for item in costs)

    logs = ok_data(client.get("/api/logs", params={"table_name": "cost", "record_id": cost_id}))["items"]
    assert any(item["action"] == "delete" for item in logs)


def test_A9_时间段筛选只统计区间内成本(client):
    project_id = STATE["project_id"]
    ok_data(
        client.post(
            "/api/costs/misc",
            json={"project_id": project_id, "category": "罚款", "amount": 800, "biz_date": "2026-02-15"},
        )
    )
    ok_data(
        client.post(
            "/api/costs/misc",
            json={"project_id": project_id, "category": "招待费", "amount": 1200, "biz_date": "2026-06-10"},
        )
    )

    ranged = ok_data(
        client.get(
            f"/api/projects/{project_id}/overview",
            params={"start_date": "2026-03-01", "end_date": "2026-06-30"},
        )
    )["metrics"]["actual_cost"]
    full = ok_data(api_overview(client, project_id))["metrics"]["actual_cost"]

    # 2/15 的 800 落在区间外
    assert round(float(full) - float(ranged), 2) == 800.00

    ranged_costs = ok_data(
        client.get(
            "/api/costs",
            params={"project_id": project_id, "start_date": "2026-03-01", "end_date": "2026-06-30", "page_size": 200},
        )
    )["items"]
    assert ranged_costs and all("2026-03-01" <= item["biz_date"] <= "2026-06-30" for item in ranged_costs)


def test_A10_导出11个Sheet且数据与页面一致(client):
    project_id = STATE["project_id"]
    response = client.get("/api/export/excel", params={"project_id": project_id})
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == [
        "项目盈亏",
        "资金流水",
        "人工",
        "材料",
        "机械",
        "杂费",
        "垫付",
        "应收",
        "应付",
        "库存",
        "人工结算",
    ]

    # 项目盈亏 Sheet 与页面数字一致
    page = ok_data(api_overview(client, project_id))["metrics"]
    sheet = workbook["项目盈亏"]
    assert sheet["F2"].value == float(page["actual_cost"])
    assert sheet["E2"].value == float(page["expect_profit"])
    assert sheet["G2"].value == float(page["settle_profit"])

    # 成本 Sheet 行数与列表一致
    misc_rows = len(ok_data(client.get("/api/costs", params={"project_id": project_id, "type": "misc"}))["items"])
    assert workbook["杂费"].max_row == misc_rows + 1
    material_rows = len(
        ok_data(client.get("/api/costs", params={"project_id": project_id, "type": "material"}))["items"]
    )
    assert workbook["材料"].max_row == material_rows + 1
