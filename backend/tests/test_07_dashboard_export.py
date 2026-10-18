"""看板、报表与 Excel 导出（FR-8、FR-9.2）。"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from helpers import add_funds, add_labor, add_material, add_misc, add_receivable, money_eq, ok_data


def test_项目盈亏总表含合计行(api, project):
    add_labor(api, project, "张三", "5", "300", "2026-05-01")
    add_material(api, project, quantity="10", unit_price="100", biz_date="2026-05-02")
    add_funds(api, project, direction="in", amount="5000", biz_date="2026-05-03")
    add_receivable(api, project, amount="8000", biz_date="2026-05-04")

    data = ok_data(api.get("/api/dashboard/projects"))
    row = next(item for item in data["items"] if item["project_id"] == project)
    money_eq(row["client_budget"], "4000000.00")
    money_eq(row["cost_budget"], "3500000.00")
    money_eq(row["expect_profit"], "500000.00")
    money_eq(row["actual_cost"], "2500.00")
    money_eq(row["settle_profit"], "3997500.00")
    money_eq(row["budget_left"], "3497500.00")
    money_eq(row["received"], "5000.00")
    money_eq(row["unreceived"], "8000.00")
    money_eq(row["paid"], "0.00")
    # 应付未结：人工 1500 + 材料 1000（两笔成本各自动生成一张应付单，尚未付款冲销）
    money_eq(row["unpaid"], "2500.00")
    assert row["name"] and row["status"] == "ongoing"

    # 合计行 = 所有项目之和
    total = data["total"]
    money_eq(total["actual_cost"], data["items"][0]["actual_cost"]) if len(data["items"]) == 1 else None
    expected_actual = sum(float(item["actual_cost"]) for item in data["items"])
    assert float(total["actual_cost"]) == expected_actual
    assert float(total["expect_profit"]) == sum(float(item["expect_profit"]) for item in data["items"])
    assert data["project_count"] == len(data["items"])


def test_时间段筛选改变实际成本但不改预计利润(api, project):
    data = ok_data(api.get("/api/dashboard/projects", params={"start_date": "2026-05-01", "end_date": "2026-05-01"}))
    row = next(item for item in data["items"] if item["project_id"] == project)
    money_eq(row["actual_cost"], "1500.00")  # 只算 5/1 的人工
    money_eq(row["expect_profit"], "500000.00")  # 预计利润与时间无关


def test_最近动态(api, project):
    data = ok_data(api.get("/api/dashboard/recent", params={"limit": 6}))
    assert 1 <= len(data["items"]) <= 6
    kinds = {item["kind"] for item in data["items"]}
    assert kinds & {"cost", "funds"}
    assert all("created_at" in item for item in data["items"])
    # 按创建时间倒序
    times = [item["created_at"] for item in data["items"]]
    assert times == sorted(times, reverse=True)


def test_按月成本趋势(api, project):
    add_misc(api, project, "伙食费", "100", "2026-06-15")
    data = ok_data(api.get("/api/dashboard/cost-trend", params={"project_id": project}))
    months = [item["month"] for item in data["items"]]
    assert months == ["2026-05", "2026-06"]
    may = data["items"][0]
    money_eq(may["total"], "2500.00")
    assert set(may["by_type"].keys()) == {"labor", "material", "machine", "misc", "advance"}
    money_eq(may["by_type"]["labor"], "1500.00")
    money_eq(may["by_type"]["material"], "1000.00")
    assert "人工" in may["by_type_label"]


def test_导出全部项目(api, project):
    response = api.get("/api/export/excel")
    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    assert "filename*=UTF-8''" in response.headers["content-disposition"]

    workbook = load_workbook(io.BytesIO(response.content))
    assert len(workbook.sheetnames) == 11
    assert workbook["人工"].max_row == 2  # 1 条工天
    assert workbook["材料"].max_row == 2
    assert workbook["资金流水"].max_row == 2
    assert workbook["应收"].max_row == 2
    assert workbook["机械"].max_row == 1  # 无数据：只有表头
    assert workbook["人工结算"].max_row == 1

    labor_sheet = workbook["人工"]
    assert [cell.value for cell in labor_sheet[1]] == ["日期", "项目", "姓名", "工日数", "日工价", "金额", "备注"]
    assert labor_sheet["A1"].font.bold is True
    assert labor_sheet.freeze_panes == "A2"
    assert labor_sheet["F2"].value == 1500
    assert labor_sheet["F2"].number_format == "#,##0.00"
    assert labor_sheet["D2"].number_format == "#,##0.00"

    funds_sheet = workbook["资金流水"]
    assert [cell.value for cell in funds_sheet[1]] == ["日期", "项目", "收付", "类型", "金额", "对方", "支付方式", "备注"]
    assert funds_sheet["C2"].value == "收"
    assert funds_sheet["G2"].value == "转账"


def test_导出按项目与时间段过滤(api, project):
    response = api.get(
        "/api/export/excel",
        params={"project_id": project, "start_date": "2026-05-01", "end_date": "2026-05-01"},
    )
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook["人工"].max_row == 2
    assert workbook["材料"].max_row == 1  # 材料在 5/2，被过滤掉

    err = api.get("/api/export/excel", params={"start_date": "2026-06-01", "end_date": "2026-05-01"})
    assert err.status_code == 400


def test_导出包含人工结算(api, project):
    add_labor(api, project, "结算工人", "2", "100", "2026-05-20")
    ok_data(api.post("/api/labor/settle", json={"project_id": project, "settle_date": "2026-07-01"}))
    response = api.get("/api/export/excel", params={"project_id": project})
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["人工结算"]

    settles = ok_data(api.get("/api/labor/settles", params={"project_id": project, "page_size": 200}))["items"]
    assert sheet.max_row == len(settles) + 1

    row_index = next(
        row for row in range(2, sheet.max_row + 1) if sheet.cell(row=row, column=2).value == "结算工人"
    )
    assert sheet.cell(row=row_index, column=3).value == 2  # 总工天
    assert sheet.cell(row=row_index, column=4).value == 200  # 应发金额
    assert sheet.cell(row=row_index, column=7).value == "未发"
