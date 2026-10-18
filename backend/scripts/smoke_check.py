"""对已启动的服务跑一遍关键业务验收（默认 http://127.0.0.1:8000）。

用法：
    .venv\\Scripts\\python.exe -m uvicorn app.main:app --port 8000
    .venv\\Scripts\\python.exe scripts\\smoke_check.py [base_url]
"""

from __future__ import annotations

import io
import os
import sys
from pathlib import Path

import httpx
from openpyxl import load_workbook

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8000"
PASSWORD = "123456"

results: list[tuple[bool, str, str]] = []


def check(name: str, passed: bool, detail: str = "") -> None:
    results.append((bool(passed), name, detail))
    print(f"{'[OK]  ' if passed else '[FAIL]'} {name}{(' → ' + detail) if detail else ''}")


def data_of(response: httpx.Response) -> dict:
    payload = response.json()
    assert payload["code"] == 0, payload
    return payload["data"]


def main() -> int:
    with httpx.Client(base_url=BASE, timeout=60) as client:
        # 0. 健康检查 / 接口文档
        health = client.get("/api/health")
        check("GET /api/health 可用", health.status_code == 200 and health.json()["code"] == 0)

        docs = client.get("/docs")
        check("GET /docs 可访问", docs.status_code == 200 and "swagger" in docs.text.lower())

        openapi = client.get("/openapi.json")
        paths = openapi.json().get("paths", {}) if openapi.status_code == 200 else {}
        check(
            "GET /openapi.json 可访问且接口齐全",
            openapi.status_code == 200 and len(paths) >= 30,
            f"{len(paths)} 个路径",
        )

        # 1. 未登录 401
        anon = httpx.get(f"{BASE}/api/projects", timeout=10)
        check("未登录访问 /api/projects 返回 401", anon.status_code == 401, f"code={anon.json()['code']}")

        # 2. 登录
        login = client.post("/api/auth/login", json={"password": PASSWORD})
        login_data = data_of(login)
        check(
            "123456 登录成功且要求设置密保",
            login_data["logged_in"] and login_data["security_ready"] is False
            and login_data["is_initial_password"] is True,
            f"security_ready={login_data['security_ready']}",
        )
        check("会话 Cookie 为 HttpOnly", "HttpOnly" in login.headers.get("set-cookie", ""))

        # 3. 建项目 400 万 / 350 万
        created = client.post(
            "/api/projects",
            json={"name": f"冒烟项目-{os.getpid()}", "client_budget": 4000000, "cost_budget": 3500000},
        )
        project = data_of(created)
        project_id = project["id"]
        check(
            "建项目 4000000/3500000 → expect_profit = 500000.00",
            project["expect_profit"] == 500000.0 and '"expect_profit":500000.00' in created.text,
            f"expect_profit={project['expect_profit']}（JSON 原文含 500000.00）",
        )

        # 4. 材料采购 100 吨 × 400
        material = data_of(
            client.post(
                "/api/costs/material",
                json={
                    "project_id": project_id,
                    "material_name": "水泥",
                    "unit": "吨",
                    "quantity": 100,
                    "unit_price": 400,
                    "supplier": "某某建材",
                    "biz_date": "2026-05-10",
                },
            )
        )
        cost_after_material = data_of(client.get(f"/api/projects/{project_id}/overview"))["metrics"]["actual_cost"]
        balance = data_of(client.get("/api/inventory", params={"project_id": project_id}))["items"][0]["balance"]
        check(
            "录材料 100 吨 × 400 → 成本 +40000 且库存余量 100",
            cost_after_material == 40000.0 and balance == 100.0,
            f"actual_cost={cost_after_material}, balance={balance}",
        )

        # 5. 领用 80 吨
        out = data_of(
            client.post(
                "/api/inventory/out",
                json={
                    "project_id": project_id,
                    "material_name": "水泥",
                    "unit": "吨",
                    "quantity": 80,
                    "biz_date": "2026-05-20",
                },
            )
        )
        cost_after_out = data_of(client.get(f"/api/projects/{project_id}/overview"))["metrics"]["actual_cost"]
        check(
            "领用 80 → 余量 20 且成本不变（仍 40000）",
            out["balance"] == 20.0 and out["enough"] is True and cost_after_out == cost_after_material,
            f"balance={out['balance']}, enough={out['enough']}, actual_cost={cost_after_out}",
        )

        # 6. 录工天 5 × 300
        labor = data_of(
            client.post(
                "/api/costs/labor",
                json={
                    "project_id": project_id,
                    "worker_name": "张三",
                    "biz_date": "2026-05-12",
                    "work_days": 5,
                    "day_rate": 300,
                },
            )
        )
        cost_after_labor = data_of(client.get(f"/api/projects/{project_id}/overview"))["metrics"]["actual_cost"]
        preview = data_of(client.get("/api/labor/preview", params={"project_id": project_id}))
        settled = data_of(client.post("/api/labor/settle", json={"project_id": project_id, "settle_date": "2026-07-01"}))
        cost_after_settle = data_of(client.get(f"/api/projects/{project_id}/overview"))["metrics"]["actual_cost"]
        check(
            "录工天 5 × 300 → 成本 +1500 且结算后成本仍是 41500",
            labor["amount"] == 1500.0
            and cost_after_labor == 41500.0
            and preview["items"][0]["amount_due"] == 1500.0
            and len(settled["items"]) == 1
            and cost_after_settle == cost_after_labor,
            f"labor={labor['amount']}, actual_cost={cost_after_labor}→{cost_after_settle}, "
            f"amount_due={preview['items'][0]['amount_due']}",
        )

        # 7. 应收 50 万 + 收款 30 万 + 冲销
        settle_profit_before = data_of(client.get(f"/api/projects/{project_id}/overview"))["metrics"]["settle_profit"]
        arap_id = data_of(
            client.post(
                "/api/arap",
                json={
                    "project_id": project_id,
                    "type": "receivable",
                    "biz_type": "progress",
                    "amount": 500000,
                    "biz_date": "2026-06-01",
                    "partner": "某某置业有限公司",
                },
            )
        )["id"]
        flow_id = data_of(
            client.post(
                "/api/funds",
                json={
                    "project_id": project_id,
                    "direction": "in",
                    "biz_type": "progress",
                    "amount": 300000,
                    "biz_date": "2026-06-20",
                    "partner": "某某置业有限公司",
                    "pay_method": "transfer",
                },
            )
        )["id"]
        offset = data_of(
            client.post(
                "/api/arap/offset",
                json={
                    "funds_flow_id": flow_id,
                    "offsets": [{"ar_ap_id": arap_id, "amount": 300000}],
                    "offset_date": "2026-06-20",
                },
            )
        )
        detail = data_of(client.get(f"/api/arap/{arap_id}"))
        settle_profit_after = data_of(client.get(f"/api/projects/{project_id}/overview"))["metrics"]["settle_profit"]
        row = next(
            item
            for item in data_of(client.get("/api/dashboard/projects"))["items"]
            if item["project_id"] == project_id
        )
        check(
            "应收 500000 收 300000 冲销 → 未收 200000 且结算利润不变",
            detail["unsettled"] == 200000.0
            and detail["status"] == "partial"
            and settle_profit_after == settle_profit_before
            and row["unreceived"] == 200000.0,
            f"unsettled={detail['unsettled']}, status={detail['status']}, "
            f"settle_profit={settle_profit_before}→{settle_profit_after}",
        )

        # 8. 撤销冲销
        undo = data_of(client.delete(f"/api/arap/offset/{offset['offset_ids'][0]}"))
        check(
            "撤销冲销 → 未收回到 500000 且状态回 open",
            undo["unsettled"] == 500000.0 and undo["status"] == "open",
            f"unsettled={undo['unsettled']}, status={undo['status']}",
        )

        # 9. 变更日志
        logs = data_of(client.get("/api/logs", params={"table_name": "cost", "page_size": 5}))
        check("变更日志有记录", logs["total"] >= 2, f"total={logs['total']}")

        # 10. 导出 Excel
        export = client.get("/api/export/excel", params={"project_id": project_id})
        workbook = load_workbook(io.BytesIO(export.content))
        expected_names = [
            "项目盈亏", "资金流水", "人工", "材料", "机械", "杂费", "垫付", "应收", "应付", "库存", "人工结算",
        ]
        check(
            "导出 Excel：11 个 Sheet、顺序正确、附件下载头正确",
            export.status_code == 200
            and workbook.sheetnames == expected_names
            and "attachment" in export.headers.get("content-disposition", ""),
            f"{len(workbook.sheetnames)} 个 Sheet，{len(export.content)} 字节",
        )

        # 11. 备份
        backup = data_of(client.post("/api/system/backup"))
        check("备份数据库文件", Path(backup["backup_path"]).exists(), backup["backup_path"])

    passed = sum(1 for item in results if item[0])
    failed = len(results) - passed
    print(f"\n冒烟结果：{passed} 项通过 / {failed} 项失败")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
