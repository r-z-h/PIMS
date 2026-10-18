"""生成《AI 项目实操手册》DOCX 的渲染器（内容分块在 manual_content_*.py）。

为什么拆文件：整本手册内容量大，分成几块更好维护，也方便以后单独改某一篇。
"""

from __future__ import annotations

import importlib
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

_PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(_PROJECT_ROOT, "docs", "07-AI项目实操手册.docx")

# 中文正文/标题字体（Word 里没有的话会自动回退，但显式设置更稳）
FONT_CN = "微软雅黑"
FONT_MONO = "Consolas"

BLUE = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xA8, 0x07, 0x07)


def _set_cn_font(style, name: str = FONT_CN) -> None:
    """中文字体必须同时设置 ascii/hAnsi 与 eastAsia，否则中文不生效。"""
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    _set_cn_font(normal)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.35

    for name, size, color in (
        ("Heading 1", 18, BLUE),
        ("Heading 2", 14, BLUE),
        ("Heading 3", 12, RGBColor(0x2E, 0x2E, 0x2E)),
        ("Heading 4", 11, GREY),
    ):
        style = doc.styles[name]
        _set_cn_font(style)
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(6)


def add_page_number_footer(section) -> None:
    footer = section.footer
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    run.font.size = Pt(9)
    run.font.color.rgb = GREY


def h(doc: Document, text: str, level: int):
    para = doc.add_heading(level=level)
    _add_rich(para, text)
    return para


def p(doc: Document, text: str, *, bold_prefix: str | None = None):
    para = doc.add_paragraph()
    if bold_prefix:
        run = para.add_run(bold_prefix)
        run.bold = True
    _add_rich(para, text)
    return para


def bullets(doc: Document, items: list[str]):
    first = None
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        _add_rich(para, item)
        first = first or para
    return first


def numbered(doc: Document, items: list[str]):
    first = None
    for item in items:
        para = doc.add_paragraph(style="List Number")
        _add_rich(para, item)
        first = first or para
    return first


def _shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


# 正文可用宽度（A4 21cm − 左右边距各 2.4cm ≈ 16.2cm），表格总宽不许超过它
TEXT_WIDTH_CM = 16.0


def _break_paragraph(doc: Document) -> None:
    """让**下一个块**从新页开始：插入一个 2 磅段落并设 page_break_before。

    比 doc.add_page_break() 安全：后者插入的是「分页符段落」，
    如果上一段刚好填满一页，那个分页符段落自己就落到了新页上，于是多出一整页空白。
    """
    para = doc.add_paragraph()
    para.paragraph_format.page_break_before = True
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1.0
    run = para.add_run(" ")
    run.font.size = Pt(2)

def _add_rich(paragraph, text: str, *, size: float | None = None, bold: bool = False):
    """把 `**加粗**` 转成真加粗（Word 不认 Markdown，直接写会显示成星号）。"""
    for index, part in enumerate(text.split("**")):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = bold or (index % 2 == 1)
        if size is not None:
            run.font.size = Pt(size)
    return paragraph

def _tiny_spacer(doc: Document) -> None:
    """表格/代码块后面的间隔段落：压到 2 磅，避免它单独撑出一页（会显示成空白页）。"""
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1.0
    run = para.add_run(" ")
    run.font.size = Pt(2)

def table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = ""
        _add_rich(hdr[i].paragraphs[0], text, size=9.5, bold=True)
        _shade(hdr[i], "DCE6F1")
    for row in rows:
        cells = tbl.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            _add_rich(cells[i].paragraphs[0], text, size=9.5)
    # 列宽：等比缩放到正文宽度以内，避免溢出到页边距（Word 里会跑到纸外）
    if widths:
        total = sum(widths)
        scale = (TEXT_WIDTH_CM / total) if total > TEXT_WIDTH_CM else 1.0
        scaled = [w * scale for w in widths]
        for row in tbl.rows:
            for i, width in enumerate(scaled):
                row.cells[i].width = Cm(round(width, 2))

    # 跨页时表头重复；单行不要被拆成两半
    header_tr_pr = tbl.rows[0]._tr.get_or_add_trPr()
    if header_tr_pr.find(qn("w:tblHeader")) is None:
        header_tr_pr.append(OxmlElement("w:tblHeader"))
    for row in tbl.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))

    _tiny_spacer(doc)


def code(doc: Document, text: str) -> None:
    for line in text.split("\n"):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.left_indent = Cm(0.6)
        para.paragraph_format.line_spacing = 1.0
        run = para.add_run(line.replace("**", "") if line else " ")
        run.font.name = FONT_MONO
        run.font.size = Pt(9)
        rpr = run._r.get_or_add_rPr()
        rfonts = OxmlElement("w:rFonts")
        rfonts.set(qn("w:ascii"), FONT_MONO)
        rfonts.set(qn("w:hAnsi"), FONT_MONO)
        rfonts.set(qn("w:eastAsia"), FONT_MONO)
        rpr.append(rfonts)
        _shade_run = OxmlElement("w:shd")
        _shade_run.set(qn("w:val"), "clear")
        _shade_run.set(qn("w:color"), "auto")
        _shade_run.set(qn("w:fill"), "F2F2F2")
        rpr.append(_shade_run)
    _tiny_spacer(doc)


def callout(doc: Document, text: str, kind: str = "note") -> None:
    color = {"note": "FFF4E5", "warn": "FDE7E9", "ok": "E8F5E9"}.get(kind, "FFF4E5")
    label = {"note": "提示", "warn": "注意", "ok": "要点"}.get(kind, "提示")
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    cell = tbl.rows[0].cells[0]
    _shade(cell, color)
    cell.text = ""
    para = cell.paragraphs[0]
    run = para.add_run(f"【{label}】")
    run.bold = True
    run.font.size = Pt(10)
    _add_rich(para, text, size=10)
    _tiny_spacer(doc)


def pagebreak(doc: Document) -> None:
    doc.add_page_break()


def render(blocks: list[tuple]) -> Document:
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.0)
    add_page_number_footer(section)

    # 分页策略：不插入「分页符段落」（那会造出空白页），
    # 而是标记 pending_break，让**下一个块自己从新页开始**。
    pending_break = False
    has_content = False

    def apply_break(paragraph) -> None:
        if paragraph is not None:
            paragraph.paragraph_format.page_break_before = True

    for block in blocks:
        kind = block[0]
        if kind == "pagebreak":
            if has_content:
                pending_break = True
            continue

        before = pending_break
        pending_break = False
        made = None

        if kind == "h1":
            made = h(doc, block[1], 1)
        elif kind == "h2":
            made = h(doc, block[1], 2)
        elif kind == "h3":
            made = h(doc, block[1], 3)
        elif kind == "p":
            made = p(doc, block[1], bold_prefix=(block[2] if len(block) > 2 else None))
        elif kind == "bullets":
            made = bullets(doc, block[1])
        elif kind == "numbered":
            made = numbered(doc, block[1])
        elif kind == "table":
            _, headers, rows, widths = block
            # 表格不能设 page_break_before，只能在它前面放一个分页段落
            if before:
                _break_paragraph(doc)
            table(doc, headers, rows, widths)
            made = None
        elif kind == "code":
            if before:
                _break_paragraph(doc)
            code(doc, block[1])
            made = None
        elif kind == "callout":
            if before:
                _break_paragraph(doc)
            callout(doc, block[1], block[2] if len(block) > 2 else "note")
            made = None
        elif kind == "title":
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(block[1])
            run.bold = True
            run.font.size = Pt(26)
            run.font.color.rgb = BLUE
            made = para
        elif kind == "subtitle":
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(block[1])
            run.font.size = Pt(12)
            run.font.color.rgb = GREY
            made = para
        elif kind == "center":
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(block[1])
            run.font.size = Pt(10)
            run.font.color.rgb = GREY
            made = para
        else:  # pragma: no cover
            raise ValueError(f"未知块类型：{kind}")

        if before:
            apply_break(made)
        has_content = True
    return doc


def drop_spacers_before_breaks(doc: Document) -> int:
    """删掉「紧挨着分页标题」的空段落。

    这些空段落本来是表格/代码块后的间隔。如果它恰好落在页面最底部，
    就会单独占掉一整页，看起来就是一张空白页。
    """
    removed = 0
    paras = doc.paragraphs
    for i in range(len(paras) - 1):
        cur = paras[i]
        if cur.text.strip():
            continue
        if paras[i + 1].paragraph_format.page_break_before:
            element = cur._element
            element.getparent().remove(element)
            removed += 1
    return removed


def main() -> None:
    blocks: list[tuple] = []
    for module_name in (
        "manual_content_0_cover",
        "manual_content_1_basics",
        "manual_content_2_glossary",
        "manual_content_3_tools",
        "manual_content_4_choices",
        "manual_content_5_deploy",
        "manual_content_6_review",
        "manual_content_7_aiwork",
        "manual_content_8_appendix",
    ):
        module = importlib.import_module(module_name)
        blocks.extend(module.BLOCKS)

    # 去掉「连续两个分页符」造成的空白页（各模块首尾都带 pagebreak）
    cleaned: list[tuple] = []
    for block in blocks:
        if block[0] == "pagebreak" and cleaned and cleaned[-1][0] == "pagebreak":
            continue
        cleaned.append(block)

    doc = render(cleaned)
    removed = drop_spacers_before_breaks(doc)
    doc.save(OUT)
    print("已生成：", OUT)
    print("块数：", len(cleaned), "（去掉连续分页符前：", len(blocks), "）")
    print("删掉分页前的空段落：", removed)


if __name__ == "__main__":
    main()
